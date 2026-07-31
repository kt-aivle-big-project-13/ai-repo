"""고영향 AI 사전진단 Word 보고서 테스트."""

from pathlib import Path

from docx import Document
from docx.shared import Cm

from app.schemas.high_impact_report import HighImpactReportRequest
from app.services.high_impact_report_docx import (
    render_high_impact_report_to_docx,
)


def _gate_answer(
    question_code: str,
    answer: bool,
) -> dict:
    return {
        "question_code": question_code,
        "question_text": f"{question_code} 정성 진단 문항",
        "stage": "QUALITATIVE",
        "group": "GATE",
        "answer": answer,
        "weight": 0,
        "score": 0,
    }


def _quantitative_answer(
    question_code: str,
    group: str,
    answer: bool,
) -> dict:
    weight = 2 if group == "A" else 1

    return {
        "question_code": question_code,
        "question_text": f"{question_code} 정량 진단 문항",
        "stage": "QUANTITATIVE",
        "group": group,
        "answer": answer,
        "weight": weight,
        "score": weight if answer else 0,
    }


def _quantitative_request() -> HighImpactReportRequest:
    return HighImpactReportRequest(
        audit_id=8,
        assessment_id=11,
        audit_name="신용평가 모델 감사",
        model_name="신용평가 모델",
        model_version="1.0",
        assessed_at="2026-07-31T10:00:00+09:00",
        condition_met=False,
        group_a_score=4,
        group_b_score=0,
        total_score=4,
        result="HIGH_IMPACT",
        answers=[
            _gate_answer("GATE_01", False),
            _gate_answer("GATE_02", False),
            _quantitative_answer("A_01", "A", True),
            _quantitative_answer("A_02", "A", True),
            _quantitative_answer("A_03", "A", False),
            _quantitative_answer("B_01", "B", False),
            _quantitative_answer("B_02", "B", False),
            _quantitative_answer("B_03", "B", False),
        ],
    )


def _qualitative_request() -> HighImpactReportRequest:
    return HighImpactReportRequest(
        audit_id=8,
        assessment_id=10,
        audit_name="신용평가 모델 감사",
        model_name="신용평가 모델",
        model_version=None,
        assessed_at="2026-07-31T10:00:00+09:00",
        condition_met=True,
        group_a_score=0,
        group_b_score=0,
        total_score=0,
        result="HIGH_IMPACT",
        answers=[
            _gate_answer("GATE_01", True),
            _gate_answer("GATE_02", False),
        ],
    )


def _document_text(path: Path) -> str:
    document = Document(path)

    paragraph_text = [
        paragraph.text
        for paragraph in document.paragraphs
    ]
    table_text = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]

    return "\n".join(paragraph_text + table_text)


def test_creates_quantitative_word_report(
    tmp_path: Path,
) -> None:
    docx_path = tmp_path / "report.docx"

    result = render_high_impact_report_to_docx(
        request=_quantitative_request(),
        decision_reason=(
            "정량 배점 결과 총 4점으로 고영향 AI로 "
            "판정되었습니다."
        ),
        docx_path=docx_path,
    )

    assert result == docx_path.resolve()
    assert docx_path.stat().st_size > 1000
    assert docx_path.read_bytes().startswith(b"PK")

    document = Document(docx_path)
    section = document.sections[0]

    assert abs(section.page_width - Cm(21.0)) <= Cm(0.01)
    assert abs(section.page_height - Cm(29.7)) <= Cm(0.01)

    text = _document_text(docx_path)

    assert "고영향 AI 사전진단 보고서" in text
    assert "정량 배점 응답" in text
    assert "A_01 정량 진단 문항" in text
    assert "합산 점수" in text
    assert "4점" in text
    assert "신용평가 모델 감사" in text


def test_qualitative_word_report_omits_quantitative_section(
    tmp_path: Path,
) -> None:
    docx_path = tmp_path / "report.docx"

    render_high_impact_report_to_docx(
        request=_qualitative_request(),
        decision_reason=(
            "정성 게이트 문항 중 1개 문항에 ‘예’로 "
            "응답하여 고영향 AI로 판정되었습니다."
        ),
        docx_path=docx_path,
    )

    text = _document_text(docx_path)

    assert "정성 게이트 응답" in text
    assert "GATE_01 정성 진단 문항" in text
    assert "정량 배점 응답" not in text
    assert "미기재" in text