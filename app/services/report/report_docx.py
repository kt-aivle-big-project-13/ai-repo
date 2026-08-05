"""설명가능성 리포트 데이터를 Word 문서로 변환한다.

용지·폰트·표·그림 같은 공통 처리는 `docx_common` 에 있고, 여기서는 설명가능성
목차와 본문 구성만 담당한다.
"""

from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.schemas.explainability.shap import ShapReport
from app.services.docx_common import (
    DocxGenerationError,
    add_figures as _add_figures,
    add_heading as _add_heading,
    add_narrative as _add_narrative,
    add_table as _add_table,
    configure_document as _configure_document,
    decode_figure as _decode_figure,
    format_list as _format_list,
    format_number as _format_number,
    set_run_font as _set_run_font,
)

# 기존 import 경로(`from app.services.report.report_docx import DocxGenerationError`)를
# 유지하기 위해 재노출한다.
__all__ = ["DocxGenerationError", "render_report_to_docx"]

RELIABILITY_METRIC_CODES = frozenset(
    {
        "GLOBAL_STABILITY",
        "FIDELITY",
        "SHAP_ADDITIVITY",
        "PERMUTATION_ALIGNMENT",
    }
)


def _add_metadata(
    document: DocumentObject,
    meta: dict[str, Any],
) -> None:
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("설명가능성 감사 리포트")
    _set_run_font(title_run, size=22, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(
        f"감사 ID {meta['audit_id']} · 분석기법 SHAP · "
        f"생성 {meta['generated_at']}"
    )
    _set_run_font(subtitle_run, size=9)

    _add_table(
        document,
        ["구분", "내용"],
        [
            ["모델", meta["model_file"]],
            ["데이터", meta["data_file"]],
            ["종합 상태", meta["overall_status"]],
        ],
    )


def render_report_to_docx(
    *,
    meta: dict[str, Any],
    report: ShapReport,
    narratives: dict[str, str],
    figures: list[dict[str, str]],
    docx_path: Path,
) -> Path:
    """구조화된 SHAP 리포트 데이터로 A4 Word 문서를 생성한다."""

    destination = docx_path.resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)

    try:
        document = Document()
        _configure_document(document)
        _add_metadata(document, meta)

        _add_heading(document, "1. 감사 개요", level=1)
        _add_heading(document, "감사 목적 및 범위", level=2)
        _add_narrative(document, narratives["overview_purpose"])

        _add_heading(document, "감사 대상 모델·데이터", level=2)
        _add_table(
            document,
            ["항목", "값"],
            [
                [
                    "모델 입력변수 수",
                    report.schema_validation.model_feature_count,
                ],
                [
                    "데이터 컬럼 수",
                    report.schema_validation.dataset_column_count,
                ],
                ["데이터 행 수", report.schema_validation.row_count],
                [
                    "분석 표본 크기",
                    _format_number(
                        report.sampling.get("audit_sample_size"),
                        0,
                    ),
                ],
            ],
        )

        _add_heading(document, "감사 결과 요약", level=2)
        _add_narrative(document, narratives["results_summary"])
        _add_table(
            document,
            ["지표", "값", "기준", "상태"],
            [
                [
                    metric.metric,
                    _format_number(metric.value),
                    _format_number(metric.threshold),
                    metric.status,
                ]
                for metric in report.metrics
            ],
        )

        _add_heading(document, "2. 모델 및 데이터 정보", level=1)
        _add_table(
            document,
            ["항목", "값"],
            [
                ["모델 파일", meta["model_file"]],
                ["데이터 파일", meta["data_file"]],
                [
                    "XGBoost 버전",
                    report.manifest.get("xgboost_version", "N/A"),
                ],
                [
                    "무작위 시드",
                    _format_number(
                        report.sampling.get("random_seed"),
                        0,
                    ),
                ],
            ],
        )

        _add_heading(document, "모델 및 데이터의 한계", level=2)
        for limitation in report.limitations:
            paragraph = document.add_paragraph(style="List Bullet")
            run = paragraph.add_run(limitation)
            _set_run_font(run, size=9.5)

        _add_heading(document, "3. 분석 전 입력 스키마 검증", level=1)
        _add_table(
            document,
            ["항목", "결과"],
            [
                ["검증 상태", report.schema_validation.status],
                [
                    "누락된 모델 입력변수",
                    _format_list(
                        report.schema_validation.missing_model_features
                    ),
                ],
                [
                    "누락된 필수 컬럼",
                    _format_list(
                        report.schema_validation.missing_required_columns
                    ),
                ],
                [
                    "중복 컬럼",
                    _format_list(report.schema_validation.duplicate_columns),
                ],
                [
                    "감사 전용 컬럼",
                    _format_list(
                        report.schema_validation.audit_only_columns
                    ),
                ],
            ],
        )

        _add_heading(document, "4. 전역 설명 분석", level=1)
        _add_heading(document, "모델의 주요 판단 경향", level=2)
        _add_narrative(document, narratives["global_interpretation"])
        _add_table(
            document,
            ["순위", "변수", "평균 |SHAP|", "영향 방향", "기여 비율", "민감"],
            [
                [
                    feature.rank,
                    feature.feature,
                    _format_number(feature.mean_abs_shap),
                    feature.direction,
                    _format_number(feature.contribution_ratio, 3),
                    "예" if feature.is_sensitive else "아니오",
                ]
                for feature in report.global_importance_top
            ],
        )
        _add_figures(
            document,
            figures,
            {
                "global_shap_top20.png",
                "explainability_dashboard.png",
            },
        )

        _add_heading(document, "5. 설명 신뢰성 검증", level=1)
        _add_narrative(document, narratives["reliability_summary"])
        _add_table(
            document,
            ["지표", "값", "기준", "상태", "부가 수치"],
            [
                [
                    metric.metric,
                    _format_number(metric.value),
                    _format_number(metric.threshold),
                    metric.status,
                    ", ".join(
                        f"{key}={_format_number(value, 3)}"
                        for key, value in metric.extra.items()
                    ),
                ]
                for metric in report.metrics
                if metric.metric in RELIABILITY_METRIC_CODES
            ],
        )
        _add_figures(
            document,
            figures,
            {
                "fidelity_distribution.png",
                "permutation_vs_shap.png",
            },
        )

        _add_heading(document, "6. 증적 및 재현성", level=1)
        _add_table(
            document,
            ["항목", "값"],
            [
                ["실행 ID", report.manifest.get("run_id", "N/A")],
                [
                    "생성 시각(UTC)",
                    report.manifest.get("generated_at_utc", "N/A"),
                ],
                [
                    "XGBoost 버전",
                    report.manifest.get("xgboost_version", "N/A"),
                ],
                [
                    "무작위 시드",
                    _format_number(
                        report.sampling.get("random_seed"),
                        0,
                    ),
                ],
            ],
        )

        _add_heading(document, "7. 종합 평가", level=1)
        _add_narrative(document, narratives["overall_assessment"])
        paragraph = document.add_paragraph()
        run = paragraph.add_run(
            f"종합 상태: {meta['overall_status']}"
        )
        _set_run_font(run, size=10, bold=True)

        _add_heading(document, "부록 · 내부 평가 기준", level=1)
        _add_table(
            document,
            ["기준 항목", "값"],
            [
                [key, _format_number(value)]
                for key, value in report.thresholds.items()
            ],
        )

        footer = document.add_paragraph()
        footer_run = footer.add_run(
            "본 리포트는 설명가능성(SHAP) 검증 결과에 한하며, "
            "민감변수·대리변수의 집단별 영향은 별도의 편향진단 "
            "리포트에서 검토함."
        )
        _set_run_font(footer_run, size=8)

        document.save(destination)
    except DocxGenerationError:
        raise
    except (KeyError, OSError, ValueError) as exception:
        raise DocxGenerationError(
            "설명가능성 Word 리포트 생성에 실패했습니다."
        ) from exception

    if not destination.is_file() or destination.stat().st_size == 0:
        raise DocxGenerationError(
            "설명가능성 Word 리포트 파일이 생성되지 않았습니다."
        )

    if not destination.read_bytes().startswith(b"PK"):
        raise DocxGenerationError(
            "생성된 설명가능성 Word 리포트 형식이 올바르지 않습니다."
        )

    return destination