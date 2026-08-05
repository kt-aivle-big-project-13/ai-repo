"""편향진단 리포트 데이터를 Word 문서로 변환한다.

HTML 리포트(`templates/bias_report.html.j2`)와 같은 목차·같은 수치를 담는다.
용지·폰트·표·그림 처리는 `docx_common` 을 쓰고, 여기서는 편향진단 목차와 본문
구성만 담당한다.
"""

from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.schemas.audit import AuditRunResponse
from app.services.docx_common import (
    DocxGenerationError,
    add_bullets,
    add_figures,
    add_heading,
    add_narrative,
    add_note,
    add_table,
    configure_document,
    format_number,
    set_run_font,
)

METRIC_NOTE = (
    "difference 계열은 0에 가까울수록, Proportional Parity는 1에 가까울수록"
    "(0.8 이상 시 80% Rule) 격차가 작음. N/A는 표본 부족 등으로 산출 불가."
)

SCOPE_NOTE = (
    "본 리포트는 판정 임계값 없이 관측 값·격차를 제시함"
    "(충족/미흡 판정은 정책 영역, 별도 후속). 설명가능성(SHAP)은 범위 밖."
)

METHOD_ITEMS = (
    "지표: Demographic Parity / Equal Opportunity / Equalized Odds / "
    "Proportional Parity(80% Rule) / FPR·FDR·FOR Parity",
    "관점: favorable(승인=1·정상=1) 기준. difference 계열은 0에 가까울수록, "
    "Proportional Parity는 1에 가까울수록 격차가 작음",
    "산출: Fairlearn + 집단별 혼동행렬(TP/FP/TN/FN) 직접 계산",
    "표본: 전체 고객을 채점해 집단별로 비교(별도 샘플링 없음). "
    "최소 표본 미만 집단은 제외",
)


def _add_title(
    document: DocumentObject,
    meta: dict[str, Any],
) -> None:
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("편향진단 감사 리포트")
    set_run_font(title_run, size=22, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(
        f"감사 ID {meta['audit_id']} · {meta['audit_name']} · "
        f"분석기법 Fairlearn · 생성 {meta['generated_at']}"
    )
    set_run_font(subtitle_run, size=9)

    add_note(document, SCOPE_NOTE)


def _add_overview(
    document: DocumentObject,
    meta: dict[str, Any],
    audit: AuditRunResponse,
    narratives: dict[str, str],
) -> None:
    add_heading(document, "1. 감사 개요", level=1)
    add_heading(document, "감사 목적 및 범위", level=2)
    add_narrative(document, narratives["overview_purpose"])

    add_table(
        document,
        ["항목", "값"],
        [
            ["고객 수", meta["n_customers"]],
            ["전체 승인율", format_number(meta["approval_rate"])],
            ["보호속성", ", ".join(audit.fairness_by_attribute.keys())],
        ],
    )


def _add_model_and_criteria(
    document: DocumentObject,
    audit: AuditRunResponse,
) -> None:
    add_heading(document, "2. 모델·데이터 및 승인 기준", level=1)

    report_meta = audit.report_meta

    if report_meta:
        add_heading(document, "모델 기본정보", level=2)
        add_table(
            document,
            ["항목", "값"],
            [
                ["모델 파일", report_meta.model_file],
                ["XGBoost 버전", report_meta.xgboost_version],
                ["입력 피처 수", report_meta.n_features],
                ["범주형 피처 수", report_meta.n_categorical_features],
            ],
        )

        add_heading(document, "분석 데이터 개요", level=2)
        add_table(
            document,
            ["항목", "값"],
            [
                ["행 수", report_meta.data_n_rows],
                ["열 수", report_meta.data_n_columns],
                ["실제값 컬럼", report_meta.target_column],
                ["보호속성 컬럼", ", ".join(report_meta.protected_columns)],
            ],
        )

    add_heading(document, "승인 기준", level=2)
    add_table(
        document,
        ["항목", "값"],
        [
            ["승인 임계값", format_number(audit.threshold.value)],
            ["산출 방식", audit.threshold.method.value],
            ["기준 근거", audit.threshold.basis],
            ["보정 출처", audit.calibration_source.value],
            [
                "전체 성능",
                f"AUC {format_number(audit.performance.auc)} · "
                f"ACC {format_number(audit.performance.accuracy)}",
            ],
        ],
    )

    if report_meta and report_meta.limitations:
        add_heading(document, "모델·데이터의 한계", level=2)
        add_bullets(document, report_meta.limitations)


def _add_method(
    document: DocumentObject,
    audit: AuditRunResponse,
) -> None:
    add_heading(document, "3. 공정성 분석 방법", level=1)
    add_bullets(document, METHOD_ITEMS)

    report_meta = audit.report_meta

    if not report_meta:
        return

    add_heading(document, "분석 전 입력 스키마 검증", level=2)

    protected_in_model = ", ".join(
        f"{column}={'예' if included else '아니오'}"
        for column, included in report_meta.protected_in_model.items()
    )

    add_table(
        document,
        ["항목", "값"],
        [
            ["검증 통과", "예" if report_meta.schema_passed else "아니오"],
            ["보호속성 모델 입력 포함", protected_in_model],
        ],
    )

    if report_meta.schema_issues:
        add_table(
            document,
            ["수준", "항목", "메시지"],
            [
                [issue.level.value, issue.item, issue.message]
                for issue in report_meta.schema_issues
            ],
        )
    else:
        add_note(document, "검증 항목에서 특이사항 없음.")


def _add_group_results(
    document: DocumentObject,
    audit: AuditRunResponse,
    narratives: dict[str, str],
    figures: list[dict[str, str]],
) -> None:
    add_heading(document, "4. 집단별 결과 (기술통계)", level=1)
    add_narrative(document, narratives["group_results"])

    for attribute, fairness in audit.fairness_by_attribute.items():
        add_heading(
            document,
            f"{attribute} ({fairness.status.value})",
            level=2,
        )

        if fairness.groups:
            add_table(
                document,
                [
                    "집단",
                    "인원",
                    "승인율",
                    "실제연체율",
                    "TP",
                    "FP",
                    "TN",
                    "FN",
                    "AUC",
                ],
                [
                    [
                        group.group,
                        group.n,
                        format_number(group.approval_rate),
                        format_number(group.actual_default_rate),
                        group.tp,
                        group.fp,
                        group.tn,
                        group.fn,
                        format_number(group.auc, 3),
                    ]
                    for group in fairness.groups
                ],
            )

        if fairness.excluded_groups:
            add_note(
                document,
                f"표본 부족 제외 집단: {', '.join(fairness.excluded_groups)}",
            )

        if fairness.note:
            add_note(document, f"비고: {fairness.note}")

    add_figures(document, figures, section="groups")


def _add_metric_results(
    document: DocumentObject,
    audit: AuditRunResponse,
    narratives: dict[str, str],
    figures: list[dict[str, str]],
) -> None:
    add_heading(document, "5. 공정성 지표 결과", level=1)
    add_narrative(document, narratives["metric_results"])

    add_table(
        document,
        [
            "보호속성",
            "DPD",
            "Equal Opp",
            "Eq Odds",
            "Prop.Parity",
            "FPR",
            "FDR",
            "FOR",
        ],
        [
            [
                attribute,
                format_number(fairness.demographic_parity_difference),
                format_number(fairness.equal_opportunity_difference),
                format_number(fairness.equalized_odds_difference),
                format_number(fairness.proportional_parity_ratio),
                format_number(fairness.fpr_parity_difference),
                format_number(fairness.fdr_parity_difference),
                format_number(fairness.for_parity_difference),
            ]
            for attribute, fairness in audit.fairness_by_attribute.items()
        ],
    )

    add_note(document, METRIC_NOTE)
    add_figures(document, figures, section="metrics")


def _add_tradeoff(
    document: DocumentObject,
    narratives: dict[str, str],
    figures: list[dict[str, str]],
) -> None:
    add_heading(document, "6. 성능–공정성 트레이드오프", level=1)
    add_narrative(document, narratives["tradeoff"])
    add_figures(document, figures, section="tradeoff")


def _add_summary(
    document: DocumentObject,
    narratives: dict[str, str],
) -> None:
    add_heading(document, "7. 종합 (기술 요약)", level=1)
    add_narrative(document, narratives["overall_summary"])
    add_note(
        document,
        "판정·개선 권고·재감사 조건·법령 매핑은 이 리포트 범위 밖"
        "(정책·후속 영역).",
    )


def _add_evidence_appendix(
    document: DocumentObject,
    audit: AuditRunResponse,
) -> None:
    report_meta = audit.report_meta

    if not report_meta:
        return

    add_heading(document, "부록 · 증적 및 재현성", level=1)

    add_heading(document, "입력 식별자", level=2)

    identifiers = [
        ["모델 S3 Key", report_meta.model_s3_key or "미기록"],
        ["모델 SHA-256", report_meta.model_sha256 or "미기록"],
        ["감사데이터 S3 Key", report_meta.audit_dataset_s3_key or "미기록"],
        ["감사데이터 SHA-256", report_meta.audit_dataset_sha256 or "미기록"],
    ]

    if report_meta.validation_dataset_s3_key:
        identifiers.extend(
            [
                ["검증데이터 S3 Key", report_meta.validation_dataset_s3_key],
                [
                    "검증데이터 SHA-256",
                    report_meta.validation_dataset_sha256 or "미기록",
                ],
            ]
        )

    add_table(document, ["항목", "값"], identifiers)

    add_heading(document, "실행 환경", level=2)
    add_table(
        document,
        ["항목", "값"],
        [
            ["실행 ID", report_meta.run_id],
            ["생성 시각(UTC)", report_meta.generated_at_utc],
            ["코드 버전", report_meta.code_version or "미기록"],
            ["XGBoost 버전", report_meta.xgboost_version],
            ["Python 버전", report_meta.python_version],
        ],
    )

    # 재현성 문구는 HTML 리포트와 동일한 조건으로 제한한다.
    if report_meta.model_sha256 and report_meta.audit_dataset_sha256:
        reproducibility = (
            "위 S3 Key·콘텐츠 해시(SHA-256)로 이번 분석에 사용된 입력을 특정할 수 "
            "있음. 동일 해시의 입력과 동일 코드 버전에서 재실행 시 결과가 재현됨"
        )

        if not report_meta.code_version:
            reproducibility += (
                " (단, 코드 버전이 미기록이라 코드 동일성은 별도 확인 필요)"
            )

        add_note(document, f"{reproducibility}.")
    else:
        add_note(
            document,
            "입력 콘텐츠 해시가 기록되지 않아, 재현성은 기록된 증적 범위 "
            "내에서만 보장됨.",
        )


def _add_warning_appendix(
    document: DocumentObject,
    audit: AuditRunResponse,
) -> None:
    if not audit.warnings:
        return

    add_heading(document, "부록 · 검증 경고", level=1)
    add_table(
        document,
        ["수준", "항목", "메시지"],
        [
            [warning.level.value, warning.item, warning.message]
            for warning in audit.warnings
        ],
    )


def render_bias_report_to_docx(
    *,
    meta: dict[str, Any],
    audit: AuditRunResponse,
    narratives: dict[str, str],
    figures: list[dict[str, str]],
    docx_path: Path,
) -> Path:
    """공정성 감사 결과로 A4 Word 문서를 생성한다."""

    destination = docx_path.resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)

    try:
        document = Document()
        configure_document(document)

        _add_title(document, meta)
        _add_overview(document, meta, audit, narratives)
        _add_model_and_criteria(document, audit)
        _add_method(document, audit)
        _add_group_results(document, audit, narratives, figures)
        _add_metric_results(document, audit, narratives, figures)
        _add_tradeoff(document, narratives, figures)
        _add_summary(document, narratives)
        _add_evidence_appendix(document, audit)
        _add_warning_appendix(document, audit)

        document.save(str(destination))
    except DocxGenerationError:
        raise
    except (OSError, KeyError, AttributeError, ValueError) as exception:
        raise DocxGenerationError(
            "편향진단 리포트의 Word 문서 생성에 실패했습니다."
        ) from exception

    return destination
