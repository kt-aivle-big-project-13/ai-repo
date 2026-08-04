"""개선 권고 가이드 생성 서비스 테스트.

LLM·S3·PDF 렌더는 monkeypatch 하고, 우선순위 산정·jinja2 렌더링·Word 조판은 실제로
수행해 문서 조립을 검증한다.
"""

from pathlib import Path

import pytest
from docx import Document

import app.services.improvement_guide as guide_service
from app.schemas.improvement_guide import (
    ComplianceGap,
    ImprovementGuideRequest,
    MetricFinding,
    SelfCheckGap,
)


def _request(**overrides) -> ImprovementGuideRequest:
    values = {
        "audit_id": 88,
        "audit_name": "테스트 감사",
        "model_name": "credit_model v1",
        "compliance_gaps": [
            ComplianceGap(
                law_name="신용정보법",
                article_no="제36조의2",
                summary="자동화 평가 결과 설명·이의제기 보장",
                evidence="자율점검 '이의제기 절차' 항목(답변: 아니요) 기반 자동 매칭",
            )
        ],
        "self_check_gaps": [
            SelfCheckGap(
                item_code="OBJECTION",
                label="심사 결과 이의제기 절차 마련 여부",
            )
        ],
        "fairness_findings": [
            MetricFinding(
                attribute="CODE_GENDER",
                metric_code="DEMOGRAPHIC_PARITY",
                value=0.45,
                threshold=0.20,
                status="FAIL",
            ),
            MetricFinding(
                attribute="AGE_GROUP",
                metric_code="EQUAL_OPPORTUNITY",
                value=0.28,
                threshold=0.20,
                status="REVIEW",
            ),
        ],
        "explainability_findings": [
            MetricFinding(
                metric_code="GLOBAL_STABILITY",
                value=0.61,
                threshold=0.70,
                status="WARNING",
            )
        ],
    }
    values.update(overrides)
    return ImprovementGuideRequest(**values)


@pytest.fixture
def patched_service(monkeypatch):
    upload_sink: dict = {}
    complete_calls: list = []

    def fake_complete(prompt, system=None, **kwargs):
        complete_calls.append(prompt)
        return "생성된 개선 서술"

    monkeypatch.setattr(guide_service, "complete", fake_complete)

    def fake_render_pdf(html_path, pdf_path):
        Path(pdf_path).write_bytes(b"%PDF-test")
        return Path(pdf_path).resolve()

    monkeypatch.setattr(guide_service, "render_html_to_pdf", fake_render_pdf)

    def fake_upload(source, key):
        upload_sink[key] = Path(source).read_bytes()
        return key

    monkeypatch.setattr(guide_service, "upload_s3_object", fake_upload)

    return upload_sink, complete_calls


def test_assigns_priority_by_source_and_status():
    """우선순위는 코드가 정한다. 법령 미준수·공정성 FAIL 은 높음이다."""

    actions = guide_service.build_actions(_request())

    by_target = {action["target"]: action for action in actions}

    assert by_target["신용정보법 제36조의2"]["priority"] == "HIGH"
    assert by_target["심사 결과 이의제기 절차 마련 여부"]["priority"] == "HIGH"
    assert by_target["CODE_GENDER · DEMOGRAPHIC_PARITY"]["priority"] == "HIGH"
    assert by_target["AGE_GROUP · EQUAL_OPPORTUNITY"]["priority"] == "MEDIUM"
    assert by_target["GLOBAL_STABILITY"]["priority"] == "LOW"

    # 높음 → 중간 → 낮음 순으로 정렬된다.
    priorities = [action["priority"] for action in actions]
    assert priorities == sorted(
        priorities,
        key=lambda value: {"HIGH": 0, "MEDIUM": 1, "LOW": 2}[value],
    )

    counts = guide_service.count_priorities(actions)
    assert counts == {"HIGH": 3, "MEDIUM": 1, "LOW": 1, "TOTAL": 5}


def test_generates_three_artifacts(patched_service):
    upload_sink, complete_calls = patched_service

    result = guide_service.generate_improvement_guide(_request())

    assert result.audit_id == 88
    assert result.format == "html"
    assert result.report_s3_key.startswith("improvement-guides/88/improvement_")
    assert result.report_s3_key.endswith("/report.html")
    assert result.pdf_report_s3_key.endswith("/report.pdf")
    assert result.word_report_s3_key.endswith("/report.docx")

    prefix = result.report_s3_key.rsplit("/", maxsplit=1)[0]
    assert result.pdf_report_s3_key.startswith(f"{prefix}/")
    assert result.word_report_s3_key.startswith(f"{prefix}/")

    # 섹션별 6회 호출 (overview·regulation·fairness·explainability·follow_up·recommendations)
    assert len(complete_calls) == 6

    assert upload_sink[result.pdf_report_s3_key].startswith(b"%PDF-")
    assert upload_sink[result.word_report_s3_key].startswith(b"PK\x03\x04")

    assert result.high_priority_count == 3
    assert result.medium_priority_count == 1
    assert result.low_priority_count == 1


def test_html_contains_actions_and_scope_note(patched_service):
    upload_sink, _ = patched_service

    result = guide_service.generate_improvement_guide(_request())
    html = upload_sink[result.report_s3_key].decode("utf-8")

    assert "개선 권고 가이드" in html
    assert "생성된 개선 서술" in html

    assert "신용정보법" in html
    assert "제36조의2" in html
    assert "CODE_GENDER" in html
    assert "GLOBAL_STABILITY" in html

    # 판정을 새로 내리지 않는다는 고지
    assert "판정을 새로 내리지 않으며" in html
    assert "우선순위 기준" in html


def test_word_document_structure(patched_service, tmp_path):
    upload_sink, _ = patched_service

    result = guide_service.generate_improvement_guide(_request())

    docx_path = tmp_path / "report.docx"
    docx_path.write_bytes(upload_sink[result.word_report_s3_key])

    document = Document(str(docx_path))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    for heading in (
        "1. 개선 개요",
        "2. 개선 과제 요약",
        "3. 규제 준수 개선",
        "4. 공정성 개선",
        "5. 설명가능성 개선",
        "6. 이행 점검 항목",
        "7. 참고 권고 (노력의무)",
        "부록 · 권고 범위와 한계",
    ):
        assert heading in text, heading

    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "제36조의2" in table_text
    assert "높음" in table_text
    assert "CODE_GENDER · DEMOGRAPHIC_PARITY" in table_text

    for table in document.tables:
        header_xml = table.rows[0]._tr.xml
        assert "tblHeader" in header_xml
        assert "cantSplit" in header_xml


def test_handles_no_action_needed(patched_service):
    """조치가 필요한 항목이 하나도 없어도 가이드가 생성된다."""

    upload_sink, _ = patched_service

    result = guide_service.generate_improvement_guide(
        _request(
            compliance_gaps=[],
            self_check_gaps=[],
            fairness_findings=[],
            explainability_findings=[],
        )
    )

    assert result.high_priority_count == 0
    assert result.medium_priority_count == 0
    assert result.low_priority_count == 0

    html = upload_sink[result.report_s3_key].decode("utf-8")
    assert "조치가 필요한 과제 없음" in html
    assert "규제 준수 영역에서 조치가 필요한 항목 없음" in html
    assert "임계값을 넘은 공정성 지표 없음" in html
    assert "임계값을 넘은 설명가능성 지표 없음" in html
    assert "노력의무 영역에서 별도로 권고할 사항 없음" in html


def test_self_check_recommendations_are_separate_from_action_list(patched_service):
    """노력의무 '아니오' 항목은 위반이 아니므로 우선순위 과제 목록엔 안 들어가고,
    별도 참고 권고 섹션에만 나타난다."""

    upload_sink, _ = patched_service

    request = _request(
        self_check_recommendations=[
            SelfCheckGap(
                item_code="IA-01",
                label="서비스 제공 전 사람의 기본권에 미치는 영향을 평가하고 문서화했나요?",
            )
        ]
    )

    actions = guide_service.build_actions(request)
    assert all("IA-01" not in action["target"] for action in actions)
    assert all(
        "기본권에 미치는 영향" not in action["detail"] for action in actions
    )

    result = guide_service.generate_improvement_guide(request)

    # 노력의무 항목이 추가돼도 우선순위 과제 수(HIGH/MEDIUM/LOW)엔 영향 없다.
    assert result.high_priority_count == 3
    assert result.medium_priority_count == 1
    assert result.low_priority_count == 1

    html = upload_sink[result.report_s3_key].decode("utf-8")
    assert "7. 참고 권고 (노력의무)" in html
    assert "서비스 제공 전 사람의 기본권에 미치는 영향을 평가하고 문서화했나요?" in html
    assert "위반이 아니라 권장 사항" in html
