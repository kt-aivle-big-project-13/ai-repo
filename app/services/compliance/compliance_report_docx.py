"""규제준수 판정서를 Word 문서로 변환한다.

HTML 리포트(`templates/compliance_report.html.j2`)와 같은 목차·같은 내용을 담는다.
용지·폰트·표 처리는 `docx_common` 을 쓴다.
"""

from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.schemas.compliance.compliance_report import ComplianceReportRequest
from app.services.docx_common import (
    DocxGenerationError,
    add_heading,
    add_narrative,
    add_note,
    add_table,
    configure_document,
    format_number,
    set_run_font,
)

VERDICT_LABELS = {
    "COMPLIANT": "준수",
    "NON_COMPLIANT": "미준수",
    "PENDING": "보류",
}

SCOPE_NOTE = (
    "본 판정서의 준수 여부는 운영기관의 자가점검 응답과 그에 매핑된 법령 조항을 "
    "근거로 함. 모델의 공정성 지표로 준수 여부를 판정한 것이 아니며, 공정성 수치는 "
    "참고값으로만 병기함."
)


def _verdict(status: str) -> str:
    return VERDICT_LABELS.get(status, status)


def _add_title(
    document: DocumentObject,
    meta: dict[str, Any],
) -> None:
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("규제준수 판정서")
    set_run_font(title_run, size=22, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    model_part = f" · 모델 {meta['model_name']}" if meta.get("model_name") else ""
    subtitle_run = subtitle.add_run(
        f"감사 ID {meta['audit_id']} · {meta['audit_name']}{model_part} · "
        f"생성 {meta['generated_at']}"
    )
    set_run_font(subtitle_run, size=9)

    add_note(document, SCOPE_NOTE)


def _add_overview(
    document: DocumentObject,
    meta: dict[str, Any],
    narratives: dict[str, str],
) -> None:
    add_heading(document, "1. 판정 개요", level=1)
    add_narrative(document, narratives["overview"])

    add_table(
        document,
        ["항목", "값"],
        [
            ["감사명", meta["audit_name"]],
            ["모델", meta.get("model_name") or "미기록"],
            ["판정 근거", "자가점검 응답 기반 법령 매핑"],
            ["판정 일시", meta["generated_at"]],
        ],
    )


def _add_verdict_summary(
    document: DocumentObject,
    counts: dict[str, int],
    narratives: dict[str, str],
) -> None:
    add_heading(document, "2. 종합 판정 결과", level=1)
    add_narrative(document, narratives["verdict_summary"])

    add_table(
        document,
        ["판정", "조항 수"],
        [
            ["준수 (COMPLIANT)", counts["COMPLIANT"]],
            ["미준수 (NON_COMPLIANT)", counts["NON_COMPLIANT"]],
            ["보류 (PENDING)", counts["PENDING"]],
            ["합계", counts["TOTAL"]],
        ],
    )


def _add_self_check(
    document: DocumentObject,
    request: ComplianceReportRequest,
) -> None:
    add_heading(document, "3. 자가점검 결과", level=1)

    add_table(
        document,
        ["항목", "응답"],
        [
            [answer.label, "예" if answer.answer else "아니오"]
            for answer in request.self_check_answers
        ],
    )

    add_note(
        document,
        "자가점검 응답은 운영기관이 스스로 답한 사실관계임. 판정의 정확성은 "
        "이 응답의 정확성에 의존함.",
    )


def _add_article_verdicts(
    document: DocumentObject,
    request: ComplianceReportRequest,
) -> None:
    add_heading(document, "4. 조항별 준수 판정", level=1)

    add_table(
        document,
        ["법령", "조항", "판정", "근거"],
        [
            [
                mapping.law_name,
                mapping.article_no,
                _verdict(mapping.compliance),
                mapping.evidence,
            ]
            for mapping in request.regulation_mappings
        ],
    )

    for mapping in request.regulation_mappings:
        add_heading(
            document,
            f"{mapping.law_name} {mapping.article_no} "
            f"({_verdict(mapping.compliance)})",
            level=2,
        )

        if mapping.summary:
            add_note(document, f"요지: {mapping.summary}")

        if mapping.content:
            add_narrative(document, mapping.content)

        if mapping.effective_date or mapping.revision_date:
            revision = (
                f" · 개정일 {mapping.revision_date}"
                if mapping.revision_date
                else ""
            )
            add_note(
                document,
                f"시행일 {mapping.effective_date or '미기록'}{revision}",
            )


def _add_non_compliance(
    document: DocumentObject,
    request: ComplianceReportRequest,
    narratives: dict[str, str],
) -> None:
    add_heading(document, "5. 미준수 항목 상세", level=1)
    add_narrative(document, narratives["non_compliance"])

    non_compliant = [
        mapping
        for mapping in request.regulation_mappings
        if mapping.compliance == "NON_COMPLIANT"
    ]

    if non_compliant:
        add_table(
            document,
            ["법령", "조항", "근거"],
            [
                [mapping.law_name, mapping.article_no, mapping.evidence]
                for mapping in non_compliant
            ],
        )
    else:
        add_note(document, "미준수로 판정된 조항 없음.")

    add_note(
        document,
        "구체적 개선 방안은 별도 산출물(개선 권고 가이드)의 범위임.",
    )


def _add_audit_reference(
    document: DocumentObject,
    request: ComplianceReportRequest,
) -> None:
    reference = request.audit_reference

    if reference is None:
        return

    add_heading(document, "6. 참고 · 감사 결과 요약", level=1)
    add_note(
        document,
        "아래 수치는 판정 근거가 아님. 공정성 지표에는 준수 여부를 가르는 "
        "임계값이 정해져 있지 않아, 맥락 제공 목적으로만 병기함.",
    )

    add_table(
        document,
        ["항목", "값"],
        [
            [
                "고객 수",
                reference.n_customers
                if reference.n_customers is not None
                else "N/A",
            ],
            ["전체 승인율", format_number(reference.approval_rate)],
            ["승인 임계값", format_number(reference.threshold)],
            ["산출 방식", reference.threshold_method or "미기록"],
            ["AUC", format_number(reference.auc)],
            ["정확도", format_number(reference.accuracy)],
        ],
    )

    if reference.fairness:
        add_table(
            document,
            ["보호속성", "DPD", "Equal Opp", "Eq Odds", "Prop.Parity"],
            [
                [
                    item.attribute,
                    format_number(item.demographic_parity_difference),
                    format_number(item.equal_opportunity_difference),
                    format_number(item.equalized_odds_difference),
                    format_number(item.proportional_parity_ratio),
                ]
                for item in reference.fairness
            ],
        )


def _add_appendix(
    document: DocumentObject,
    request: ComplianceReportRequest,
    counts: dict[str, int],
    narratives: dict[str, str],
) -> None:
    add_heading(document, "부록 · 판정 방법과 한계", level=1)
    add_narrative(document, narratives["limitation"])

    add_table(
        document,
        ["항목", "값"],
        [
            ["판정 방식", "자가점검 응답 → 정적 매핑표 → 법령 조항별 준수 판정"],
            ["자가점검 항목 수", len(request.self_check_answers)],
            ["판정 대상 조항 수", counts["TOTAL"]],
        ],
    )

    add_note(
        document,
        "공정성 지표 기반 준수 판정은 판정 임계값이 정책으로 확정된 뒤에 별도로 다룸.",
    )


def render_compliance_report_to_docx(
    *,
    meta: dict[str, Any],
    request: ComplianceReportRequest,
    counts: dict[str, int],
    narratives: dict[str, str],
    docx_path: Path,
) -> Path:
    """규제준수 판정 결과로 A4 Word 문서를 생성한다."""

    destination = docx_path.resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)

    try:
        document = Document()
        configure_document(document)

        _add_title(document, meta)
        _add_overview(document, meta, narratives)
        _add_verdict_summary(document, counts, narratives)
        _add_self_check(document, request)
        _add_article_verdicts(document, request)
        _add_non_compliance(document, request, narratives)
        _add_audit_reference(document, request)
        _add_appendix(document, request, counts, narratives)

        document.save(str(destination))
    except DocxGenerationError:
        raise
    except (OSError, KeyError, AttributeError, ValueError) as exception:
        raise DocxGenerationError(
            "규제준수 판정서의 Word 문서 생성에 실패했습니다."
        ) from exception

    return destination
