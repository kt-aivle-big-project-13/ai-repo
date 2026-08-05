"""편향진단 리포트 Word 렌더러 테스트.

figure 는 실제 PNG 를 base64 로 넣어 이미지 삽입까지 확인하고, 잘못된 데이터는
`DocxGenerationError` 로 감싸지는지 본다.
"""

import base64

import pytest
from docx import Document
from docx.shared import Cm

from app.services.bias.bias_report_docx import render_bias_report_to_docx
from app.services.docx_common import DocxGenerationError
from tests.test_bias_report import _audit

# 유효한 1x1 PNG 이미지 (설명가능성 Word 테스트와 동일한 픽스처).
PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAAB"
    "CAQAAAC1HAwCAAAAC0lEQVR42mNk+A8A"
    "AQUBAScY42YAAAAASUVORK5CYII="
)

NARRATIVES = {
    "overview_purpose": "감사 목적 서술",
    "group_results": "집단별 결과 서술",
    "metric_results": "지표 결과 서술",
    "tradeoff": "트레이드오프 서술",
    "overall_summary": "종합 서술",
}


def _meta() -> dict:
    return {
        "audit_id": 42,
        "audit_name": "테스트 감사",
        "generated_at": "2026-07-31T00:00:00+00:00",
        "n_customers": 5500,
        "approval_rate": 0.84,
    }


def _figures(data_uri: str | None = None) -> list[dict[str, str]]:
    encoded = base64.b64encode(PNG_BYTES).decode("ascii")
    uri = data_uri or f"data:image/png;base64,{encoded}"

    return [
        {
            "name": "group_rates_CODE_GENDER",
            "attribute": "CODE_GENDER",
            "section": "groups",
            "title": "집단별 승인율·연체율",
            "data_uri": uri,
        },
        {
            "name": "metric_bars",
            "attribute": "CODE_GENDER",
            "section": "metrics",
            "title": "보호속성별 지표",
            "data_uri": uri,
        },
        {
            "name": "tradeoff_curve",
            "attribute": "CODE_GENDER",
            "section": "tradeoff",
            "title": "성능-공정성 트레이드오프",
            "data_uri": uri,
        },
    ]


def test_render_bias_report_to_docx_creates_word_document(tmp_path):
    docx_path = tmp_path / "report.docx"

    result = render_bias_report_to_docx(
        meta=_meta(),
        audit=_audit(),
        narratives=NARRATIVES,
        figures=_figures(),
        docx_path=docx_path,
    )

    assert result.exists()
    assert result.read_bytes().startswith(b"PK\x03\x04")

    document = Document(str(result))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    # 제목과 확정 목차 7개 장 + 부록
    assert "편향진단 감사 리포트" in text
    for heading in (
        "1. 감사 개요",
        "2. 모델·데이터 및 승인 기준",
        "3. 공정성 분석 방법",
        "4. 집단별 결과 (기술통계)",
        "5. 공정성 지표 결과",
        "6. 성능–공정성 트레이드오프",
        "7. 종합 (기술 요약)",
        "부록 · 증적 및 재현성",
    ):
        assert heading in text, heading

    # LLM 서술이 모두 들어갔는지
    for narrative in NARRATIVES.values():
        assert narrative in text, narrative

    # 표: 집단표·지표표·증적표
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "Prop.Parity" in table_text
    assert "실제연체율" in table_text
    assert "models/abc_credit_model.json" in table_text
    assert "a" * 64 in table_text

    # figure 3개가 모두 섹션별로 삽입됐는지
    assert len(document.inline_shapes) == 3

    # A4 용지 설정 (twips 로 저장되며 되읽을 때 반올림 오차가 있어 허용범위로 본다)
    section = document.sections[0]
    assert abs(section.page_width - Cm(21.0)) < 1000
    assert abs(section.page_height - Cm(29.7)) < 1000


def test_render_bias_report_to_docx_repeats_table_headers(tmp_path):
    """표가 페이지를 넘어가도 머리글이 반복되고 행이 쪼개지지 않아야 한다."""

    docx_path = tmp_path / "report.docx"

    render_bias_report_to_docx(
        meta=_meta(),
        audit=_audit(),
        narratives=NARRATIVES,
        figures=_figures(),
        docx_path=docx_path,
    )

    document = Document(str(docx_path))

    for table in document.tables:
        header_xml = table.rows[0]._tr.xml
        assert "tblHeader" in header_xml
        assert "cantSplit" in header_xml


def test_render_bias_report_to_docx_rejects_invalid_figure(tmp_path):
    docx_path = tmp_path / "report.docx"

    with pytest.raises(DocxGenerationError):
        render_bias_report_to_docx(
            meta=_meta(),
            audit=_audit(),
            narratives=NARRATIVES,
            figures=_figures(data_uri="not-a-data-uri"),
            docx_path=docx_path,
        )
