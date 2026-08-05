"""고영향 AI 사전진단 Word 보고서 렌더러."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.schemas.highimpact.high_impact_report import HighImpactReportRequest
from app.services.docx_common import (
    DocxGenerationError,
    add_heading,
    add_narrative,
    add_note,
    add_table,
    configure_document,
    set_run_font,
)

QUESTION_ORDER = (
    "GATE_01",
    "GATE_02",
    "A_01",
    "A_02",
    "A_03",
    "B_01",
    "B_02",
    "B_03",
)


def _answer_label(answer: bool) -> str:
    return "예" if answer else "아니요"


def render_high_impact_report_to_docx(
    *,
    request: HighImpactReportRequest,
    decision_reason: str,
    docx_path: Path,
) -> Path:
    """사전진단 데이터를 A4 Word 문서로 생성한다."""

    try:
        document = Document()
        configure_document(document)

        title = document.add_heading(
            "고영향 AI 사전진단 보고서",
            level=0,
        )
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_after = Pt(14)

        subtitle_run = subtitle.add_run(
            "저장된 문항별 응답과 내부 판정 기준에 따라 "
            "자동 생성된 결과서입니다."
        )
        set_run_font(subtitle_run, size=9)

        add_heading(document, "1. 사전진단 개요", level=1)
        add_table(
            document,
            ["항목", "내용"],
            [
                ["감사 ID", request.audit_id],
                ["사전진단 ID", request.assessment_id],
                ["감사명", request.audit_name],
                [
                    "진단 일시",
                    request.assessed_at.strftime(
                        "%Y-%m-%d %H:%M:%S %z"
                    ),
                ],
                ["모델명", request.model_name],
                [
                    "모델 버전",
                    request.model_version or "미기재",
                ],
            ],
        )

        add_heading(document, "2. 최종 판정", level=1)
        result_label = (
            "고영향 AI"
            if request.result == "HIGH_IMPACT"
            else "고영향 AI 미해당"
        )
        add_table(
            document,
            ["판정 항목", "결과"],
            [["고영향 AI 해당 여부", result_label]],
        )
        add_narrative(document, decision_reason)

        answer_by_code = {
            answer.question_code: answer
            for answer in request.answers
        }
        ordered_answers = [
            answer_by_code[question_code]
            for question_code in QUESTION_ORDER
            if question_code in answer_by_code
        ]

        qualitative_answers = [
            answer
            for answer in ordered_answers
            if answer.stage == "QUALITATIVE"
        ]
        quantitative_answers = [
            answer
            for answer in ordered_answers
            if answer.stage == "QUANTITATIVE"
        ]

        add_heading(document, "3. 정성 게이트 응답", level=1)
        add_table(
            document,
            ["문항 코드", "진단 문항", "응답"],
            [
                [
                    answer.question_code,
                    answer.question_text,
                    _answer_label(answer.answer),
                ]
                for answer in qualitative_answers
            ],
        )

        next_section = 4

        if quantitative_answers:
            add_heading(
                document,
                f"{next_section}. 정량 배점 응답",
                level=1,
            )
            add_table(
                document,
                [
                    "문항 코드",
                    "진단 문항",
                    "응답",
                    "배점",
                    "획득 점수",
                ],
                [
                    [
                        answer.question_code,
                        answer.question_text,
                        _answer_label(answer.answer),
                        f"{answer.weight}점",
                        f"{answer.score}점",
                    ]
                    for answer in quantitative_answers
                ],
            )
            next_section += 1

        add_heading(
            document,
            f"{next_section}. 점수 산정 결과",
            level=1,
        )
        add_table(
            document,
            ["A그룹 점수", "B그룹 점수", "합산 점수"],
            [
                [
                    f"{request.group_a_score}점",
                    f"{request.group_b_score}점",
                    f"{request.total_score}점",
                ]
            ],
        )
        next_section += 1

        add_heading(
            document,
            f"{next_section}. 판정 기준",
            level=1,
        )
        add_narrative(
            document,
            "정성 게이트 문항 중 하나 이상이 ‘예’이면 "
            "정량 배점 없이 고영향 AI로 판정합니다.",
        )
        add_narrative(
            document,
            "정성 게이트를 충족하지 않은 경우 A그룹과 "
            "B그룹의 합산 점수가 4점 이상이면 "
            "고영향 AI로 판정합니다.",
        )
        next_section += 1

        add_heading(
            document,
            f"{next_section}. 유의사항",
            level=1,
        )
        add_note(
            document,
            "본 보고서는 사용자가 제출한 사전진단 응답과 "
            "시스템 내부 판정 기준을 기반으로 자동 생성되었습니다.",
        )
        add_note(
            document,
            "실제 법적 의무와 적용 범위는 관련 법령, 감독기관 "
            "지침 및 전문가 검토 결과에 따라 달라질 수 있습니다.",
        )

        docx_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )
        document.save(docx_path)

        return docx_path.resolve()
    except DocxGenerationError:
        raise
    except (KeyError, OSError, TypeError, ValueError) as exception:
        raise DocxGenerationError(
            "고영향 AI 사전진단 Word 보고서 생성에 실패했습니다."
        ) from exception