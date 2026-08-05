"""설명가능성 리포트 Word 문서 생성 테스트."""

import base64
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.image.exceptions import (
    InvalidImageStreamError,
    UnexpectedEndOfFileError,
    UnrecognizedImageError,
)
from docx.shared import Cm
from docx.text.run import Run

from app.services.report.report_docx import (
    DocxGenerationError,
    render_report_to_docx,
)



# 유효한 1x1 PNG 이미지.
PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAAB"
    "CAQAAAC1HAwCAAAAC0lEQVR42mNk+A8A"
    "AQUBAScY42YAAAAASUVORK5CYII="
)


def _report():
    metric = SimpleNamespace(
        metric="GLOBAL_STABILITY",
        value=0.82,
        threshold=0.75,
        status="PASS",
        extra={"standard_deviation": 0.03},
    )

    sensitive_metric = SimpleNamespace(
        metric="SENSITIVE_CONTRIB",
        value=0.12,
        threshold=0.10,
        status="WARNING",
        extra={},
    )

    feature = SimpleNamespace(
        rank=1,
        feature="EXT_SOURCE_3",
        mean_abs_shap=0.1234,
        direction="RISK_DECREASE",
        contribution_ratio=0.42,
        is_sensitive=False,
    )

    schema_validation = SimpleNamespace(
        status="PASS",
        model_feature_count=50,
        dataset_column_count=52,
        row_count=1000,
        missing_model_features=[],
        missing_required_columns=[],
        duplicate_columns=[],
        audit_only_columns=["CODE_GENDER"],
    )

    return SimpleNamespace(
        schema_validation=schema_validation,
        metrics=[metric, sensitive_metric],
        global_importance_top=[feature],
        sampling={
            "audit_sample_size": 500,
            "random_seed": 42,
        },
        thresholds={
            "global_stability_threshold": 0.75,
        },
        manifest={
            "run_id": "shap_audit_test",
            "generated_at_utc": "2026-07-30T00:00:00+00:00",
            "xgboost_version": "3.1.0",
        },
        limitations=["표본 기반 분석 결과임"],
    )


def _meta() -> dict:
    return {
        "audit_id": 42,
        "generated_at": "2026-07-30T00:00:00+00:00",
        "overall_status": "WARNING",
        "model_file": "credit_model.json",
        "data_file": "audit_dataset.csv",
    }


def _narratives() -> dict[str, str]:
    return {
        "overview_purpose": "감사 목적 및 범위 설명",
        "results_summary": "감사 결과 요약 설명",
        "global_interpretation": "주요 판단 경향 설명",
        "reliability_summary": "설명 신뢰성 검증 설명",
        "overall_assessment": "종합 평가 설명",
    }


def _figures() -> list[dict[str, str]]:
    encoded = base64.b64encode(PNG_BYTES).decode("ascii")

    return [
        {
            "name": "global_shap_top20.png",
            "title": "전역 SHAP 중요도 상위 20",
            "data_uri": f"data:image/png;base64,{encoded}",
        }
    ]


def test_render_report_to_docx_creates_word_document(
    tmp_path: Path,
) -> None:
    docx_path = tmp_path / "report.docx"

    result = render_report_to_docx(
        meta=_meta(),
        report=_report(),
        narratives=_narratives(),
        figures=_figures(),
        docx_path=docx_path,
    )

    assert result == docx_path.resolve()
    assert docx_path.stat().st_size > 1000
    assert docx_path.read_bytes().startswith(b"PK")

    document = Document(docx_path)
    paragraph_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )

    assert "설명가능성 감사 리포트" in paragraph_text
    assert "감사 목적 및 범위 설명" in paragraph_text
    assert "종합 평가 설명" in paragraph_text
    assert "EXT_SOURCE_3" in table_text
    assert "GLOBAL_STABILITY" in table_text
    summary_table_text = "\n".join(
        cell.text
        for row in document.tables[2].rows
        for cell in row.cells
    )
    reliability_table_text = "\n".join(
        cell.text
        for row in document.tables[6].rows
        for cell in row.cells
    )
    assert "SENSITIVE_CONTRIB" in summary_table_text
    assert "SENSITIVE_CONTRIB" not in reliability_table_text
    assert "GLOBAL_STABILITY" in reliability_table_text
    assert len(document.tables) >= 7
    assert len(document.inline_shapes) == 1

    section = document.sections[0]
    assert abs(section.page_width - Cm(21.0)) < 1000
    assert abs(section.page_height - Cm(29.7)) < 1000


def test_render_report_to_docx_rejects_invalid_figure(
    tmp_path: Path,
) -> None:
    docx_path = tmp_path / "report.docx"
    invalid_figures = [
        {
            "name": "global_shap_top20.png",
            "title": "잘못된 이미지",
            "data_uri": "data:image/png;base64,not-base64",
        }
    ]

    with pytest.raises(
        DocxGenerationError,
        match="그래프 데이터를 읽을 수 없습니다",
    ):
        render_report_to_docx(
            meta=_meta(),
            report=_report(),
            narratives=_narratives(),
            figures=invalid_figures,
            docx_path=docx_path,
        )


@pytest.mark.parametrize(
    "image_error",
    [
        InvalidImageStreamError,
        UnexpectedEndOfFileError,
        UnrecognizedImageError,
    ],
)
def test_render_report_to_docx_wraps_corrupted_image_errors(
    tmp_path: Path,
    monkeypatch,
    image_error,
) -> None:
    docx_path = tmp_path / "report.docx"

    def fail_to_add_picture(self, *args, **kwargs):
        raise image_error()

    monkeypatch.setattr(
        Run,
        "add_picture",
        fail_to_add_picture,
    )

    with pytest.raises(
        DocxGenerationError,
        match="그래프 형식이 올바르지 않습니다",
    ):
        render_report_to_docx(
            meta=_meta(),
            report=_report(),
            narratives=_narratives(),
            figures=_figures(),
            docx_path=docx_path,
        )
