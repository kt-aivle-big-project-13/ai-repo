"""개선 권고 가이드를 Word 문서로 변환한다.

HTML 리포트(`templates/improvement_guide.html.j2`)와 같은 목차·같은 내용을 담는다.
"""

from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.schemas.improvement_guide import ImprovementGuideRequest
from app.services.docx_common import (
    DocxGenerationError,
    add_heading,
    add_narrative,
    add_note,
    add_table,
    configure_document,
    set_run_font,
)

PRIORITY_LABELS = {"HIGH": "높음", "MEDIUM": "중간", "LOW": "낮음"}

SCOPE_NOTE = (
    "본 가이드는 이미 임계값 판정을 거쳐 선별된 항목에 대한 개선 방향을 제시함. "
    "판정을 새로 내리지 않으며, 우선순위는 항목의 성격과 판정 상태에 따라 정해짐."
)

LIMITATION_NOTE = (
    "규제 준수 판단은 자가점검 응답에 기반하므로 응답의 정확성에 의존함. "
    "공정성·설명가능성 지표가 기준 이내라는 것이 문제가 없음을 보장하지는 않음. "
    "재감사 시점과 구체적 목표 수치는 이 가이드의 범위 밖임."
)


def _priority(value: str) -> str:
    return PRIORITY_LABELS.get(value, value)


def _metric_value(value: float | None) -> str:
    return "N/A" if value is None else str(value)


def _add_title(
    document: DocumentObject,
    meta: dict[str, Any],
) -> None:
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("개선 권고 가이드")
    set_run_font(title_run, size=22, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    model_part = f" · 모델 {meta['model_name']}" if meta.get("model_name") else ""
    subtitle_run = subtitle.add_run(
        f"감사 ID {meta['audit_id']} · {meta['audit_name']}{model_part} · "
        f"생성 {meta['generated_at']}"
    )
    set_run_font(subtitle_run, size=9)

    add_note(document, SCOPE_NOTE)


def _add_overview(
    document: DocumentObject,
    meta: dict[str, Any],
    counts: dict[str, int],
    narratives: dict[str, str],
) -> None:
    add_heading(document, "1. 개선 개요", level=1)
    add_narrative(document, narratives["overview"])

    add_table(
        document,
        ["항목", "값"],
        [
            ["감사명", meta["audit_name"]],
            ["모델", meta.get("model_name") or "미기록"],
            ["전체 과제 수", counts["TOTAL"]],
            ["생성 일시", meta["generated_at"]],
        ],
    )


def _add_action_summary(
    document: DocumentObject,
    actions: list[dict[str, Any]],
    counts: dict[str, int],
) -> None:
    add_heading(document, "2. 개선 과제 요약", level=1)

    add_table(
        document,
        ["우선순위", "과제 수"],
        [
            ["높음", counts["HIGH"]],
            ["중간", counts["MEDIUM"]],
            ["낮음", counts["LOW"]],
            ["합계", counts["TOTAL"]],
        ],
    )

    if actions:
        add_table(
            document,
            ["우선순위", "영역", "대상", "내용"],
            [
                [
                    _priority(action["priority"]),
                    action["area"],
                    action["target"],
                    action["detail"],
                ]
                for action in actions
            ],
        )
    else:
        add_note(document, "조치가 필요한 과제 없음.")


def _add_regulation(
    document: DocumentObject,
    request: ImprovementGuideRequest,
    narratives: dict[str, str],
) -> None:
    add_heading(document, "3. 규제 준수 개선", level=1)
    add_narrative(document, narratives["regulation"])

    if request.compliance_gaps:
        add_heading(document, "미준수 법령 조항", level=2)
        add_table(
            document,
            ["법령", "조항", "요지", "근거"],
            [
                [
                    gap.law_name,
                    gap.article_no,
                    gap.summary or "미기록",
                    gap.evidence,
                ]
                for gap in request.compliance_gaps
            ],
        )

    if request.self_check_gaps:
        add_heading(document, "자가점검 미충족 항목", level=2)
        add_table(
            document,
            ["항목"],
            [[gap.label] for gap in request.self_check_gaps],
        )

    if not request.compliance_gaps and not request.self_check_gaps:
        add_note(document, "규제 준수 영역에서 조치가 필요한 항목 없음.")


def _add_fairness(
    document: DocumentObject,
    request: ImprovementGuideRequest,
    narratives: dict[str, str],
) -> None:
    add_heading(document, "4. 공정성 개선", level=1)
    add_narrative(document, narratives["fairness"])

    if request.fairness_findings:
        add_table(
            document,
            ["보호속성", "지표", "관측값", "임계값", "상태"],
            [
                [
                    finding.attribute or "-",
                    finding.metric_code,
                    _metric_value(finding.value),
                    _metric_value(finding.threshold),
                    finding.status,
                ]
                for finding in request.fairness_findings
            ],
        )
    else:
        add_note(document, "임계값을 넘은 공정성 지표 없음.")


def _add_explainability(
    document: DocumentObject,
    request: ImprovementGuideRequest,
    narratives: dict[str, str],
) -> None:
    add_heading(document, "5. 설명가능성 개선", level=1)
    add_narrative(document, narratives["explainability"])

    if request.explainability_findings:
        add_table(
            document,
            ["지표", "관측값", "임계값", "상태"],
            [
                [
                    finding.metric_code,
                    _metric_value(finding.value),
                    _metric_value(finding.threshold),
                    finding.status,
                ]
                for finding in request.explainability_findings
            ],
        )
    else:
        add_note(document, "임계값을 넘은 설명가능성 지표 없음.")


def _add_follow_up(
    document: DocumentObject,
    narratives: dict[str, str],
) -> None:
    add_heading(document, "6. 이행 점검 항목", level=1)
    add_narrative(document, narratives["follow_up"])


def _add_recommendations(
    document: DocumentObject,
    request: ImprovementGuideRequest,
    narratives: dict[str, str],
) -> None:
    add_heading(document, "7. 참고 권고 (노력의무)", level=1)
    add_note(
        document,
        "아래 항목은 법이 \"노력하여야 한다\"고 정한 노력의무라, 위반이 아니라 권장 사항임.",
    )
    add_narrative(document, narratives["recommendations"])

    if request.self_check_recommendations:
        add_table(
            document,
            ["항목"],
            [[gap.label] for gap in request.self_check_recommendations],
        )
    else:
        add_note(document, "노력의무 영역에서 별도로 권고할 사항 없음.")


def _add_appendix(document: DocumentObject) -> None:
    add_heading(document, "부록 · 권고 범위와 한계", level=1)

    add_table(
        document,
        ["항목", "값"],
        [
            [
                "선별 방식",
                "백엔드에 저장된 임계값 판정 결과 중 조치가 필요한 항목만 포함",
            ],
            [
                "우선순위 기준",
                "법령 미준수·공정성 FAIL은 높음, 공정성 REVIEW·설명가능성 REVIEW는 "
                "중간, 설명가능성 WARNING은 낮음",
            ],
        ],
    )

    add_note(document, LIMITATION_NOTE)


def render_improvement_guide_to_docx(
    *,
    meta: dict[str, Any],
    request: ImprovementGuideRequest,
    actions: list[dict[str, Any]],
    counts: dict[str, int],
    narratives: dict[str, str],
    docx_path: Path,
) -> Path:
    """개선 과제 목록으로 A4 Word 문서를 생성한다."""

    destination = docx_path.resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)

    try:
        document = Document()
        configure_document(document)

        _add_title(document, meta)
        _add_overview(document, meta, counts, narratives)
        _add_action_summary(document, actions, counts)
        _add_regulation(document, request, narratives)
        _add_fairness(document, request, narratives)
        _add_explainability(document, request, narratives)
        _add_follow_up(document, narratives)
        _add_recommendations(document, request, narratives)
        _add_appendix(document)

        document.save(str(destination))
    except DocxGenerationError:
        raise
    except (OSError, KeyError, AttributeError, ValueError) as exception:
        raise DocxGenerationError(
            "개선 권고 가이드의 Word 문서 생성에 실패했습니다."
        ) from exception

    return destination
