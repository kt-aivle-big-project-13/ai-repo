"""규제준수 판정서 생성 서비스 테스트.

LLM·S3·PDF 렌더는 monkeypatch 하고, jinja2 렌더링과 Word 조판은 실제로 수행해
문서 조립을 검증한다.
"""

from pathlib import Path

import pytest
from docx import Document

import app.services.compliance_report as compliance_service
from app.schemas.compliance_report import (
    AuditReference,
    ComplianceReportRequest,
    FairnessReference,
    RegulationMapping,
    SelfCheckAnswer,
)


def _request(**overrides) -> ComplianceReportRequest:
    values = {
        "audit_id": 77,
        "audit_name": "테스트 감사",
        "model_name": "credit_model v1",
        "self_check_answers": [
            SelfCheckAnswer(
                item_code="NOTICE",
                label="AI 심사 사실 사전 고지 여부",
                answer=True,
            ),
            SelfCheckAnswer(
                item_code="OBJECTION",
                label="심사 결과 이의제기 절차 마련 여부",
                answer=False,
            ),
        ],
        "regulation_mappings": [
            RegulationMapping(
                law_name="신용정보법",
                article_no="제36조의2",
                content="개인신용평점 산정 결과에 대한 설명 요구권 조항 본문",
                summary="자동화 평가 결과 설명·이의제기 보장",
                compliance="NON_COMPLIANT",
                evidence="자율점검 '심사 결과 이의제기 절차 마련 여부' 항목(답변: 아니요) 기반 자동 매칭",
                effective_date="2020-08-05",
                revision_date="2024-03-01",
            ),
            RegulationMapping(
                law_name="AI 기본법",
                article_no="제27조",
                content="고영향 AI 고지 의무 조항 본문",
                summary="이용자에 대한 사전 고지",
                compliance="COMPLIANT",
                evidence="자율점검 'AI 심사 사실 사전 고지 여부' 항목(답변: 예) 기반 자동 매칭",
            ),
        ],
        "audit_reference": AuditReference(
            n_customers=5500,
            approval_rate=0.84,
            threshold=0.42,
            threshold_method="VALIDATION_DATASET",
            auc=0.755,
            accuracy=0.71,
            fairness=[
                FairnessReference(
                    attribute="CODE_GENDER",
                    demographic_parity_difference=0.08,
                    equal_opportunity_difference=0.05,
                    equalized_odds_difference=0.06,
                    proportional_parity_ratio=0.9,
                )
            ],
        ),
    }
    values.update(overrides)
    return ComplianceReportRequest(**values)


@pytest.fixture
def patched_service(monkeypatch):
    """LLM·PDF·S3 를 mock 하고 업로드된 산출물을 모아준다."""

    upload_sink: dict = {}
    complete_calls: list = []

    def fake_complete(prompt, system=None, **kwargs):
        complete_calls.append(prompt)
        return "생성된 판정 서술"

    monkeypatch.setattr(compliance_service, "complete", fake_complete)

    def fake_render_pdf(html_path, pdf_path):
        Path(pdf_path).write_bytes(b"%PDF-test")
        return Path(pdf_path).resolve()

    monkeypatch.setattr(compliance_service, "render_html_to_pdf", fake_render_pdf)

    def fake_upload(source, key):
        upload_sink[key] = Path(source).read_bytes()
        return key

    monkeypatch.setattr(compliance_service, "upload_s3_object", fake_upload)

    return upload_sink, complete_calls


def test_generates_three_artifacts(patched_service):
    upload_sink, complete_calls = patched_service

    result = compliance_service.generate_compliance_report(_request())

    assert result.audit_id == 77
    assert result.format == "html"
    assert result.report_s3_key.startswith("compliance-reports/77/compliance_")
    assert result.report_s3_key.endswith("/report.html")
    assert result.pdf_report_s3_key.endswith("/report.pdf")
    assert result.word_report_s3_key.endswith("/report.docx")

    # 세 산출물이 같은 실행(prefix)에 묶여야 백엔드가 한 건으로 저장한다.
    prefix = result.report_s3_key.rsplit("/", maxsplit=1)[0]
    assert result.pdf_report_s3_key.startswith(f"{prefix}/")
    assert result.word_report_s3_key.startswith(f"{prefix}/")

    # 섹션별 4회 호출
    assert len(complete_calls) == 4

    assert upload_sink[result.pdf_report_s3_key].startswith(b"%PDF-")
    assert upload_sink[result.word_report_s3_key].startswith(b"PK\x03\x04")


def test_counts_verdicts_without_deciding(patched_service):
    """판정은 요청에 담겨 오고, 서비스는 집계만 한다."""

    upload_sink, _ = patched_service

    result = compliance_service.generate_compliance_report(_request())

    assert result.compliant_count == 1
    assert result.non_compliant_count == 1
    assert result.pending_count == 0


def test_html_contains_verdicts_and_scope_note(patched_service):
    upload_sink, _ = patched_service

    result = compliance_service.generate_compliance_report(_request())
    html = upload_sink[result.report_s3_key].decode("utf-8")

    assert "규제준수 판정서" in html
    assert "생성된 판정 서술" in html

    # 판정 결과
    assert "신용정보법" in html
    assert "제36조의2" in html
    assert "미준수" in html
    assert "준수" in html

    # 자가점검 결과
    assert "AI 심사 사실 사전 고지 여부" in html

    # 판정 근거가 자가점검임을 밝히는 문구 (핵심 원칙)
    assert "자가점검 응답" in html
    assert "판정 근거가 아님" in html

    # 참고 감사 수치
    assert "CODE_GENDER" in html


def test_word_document_structure(patched_service, tmp_path):
    upload_sink, _ = patched_service

    result = compliance_service.generate_compliance_report(_request())

    docx_path = tmp_path / "report.docx"
    docx_path.write_bytes(upload_sink[result.word_report_s3_key])

    document = Document(str(docx_path))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    for heading in (
        "1. 판정 개요",
        "2. 종합 판정 결과",
        "3. 자가점검 결과",
        "4. 조항별 준수 판정",
        "5. 미준수 항목 상세",
        "6. 참고 · 감사 결과 요약",
        "부록 · 판정 방법과 한계",
    ):
        assert heading in text, heading

    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "제36조의2" in table_text
    assert "미준수" in table_text
    assert "CODE_GENDER" in table_text

    # 표 머리글 반복·행 분할 방지
    for table in document.tables:
        header_xml = table.rows[0]._tr.xml
        assert "tblHeader" in header_xml
        assert "cantSplit" in header_xml


def test_handles_all_compliant(patched_service):
    """미준수가 없어도 5장이 비지 않고 정상 생성된다."""

    upload_sink, _ = patched_service

    request = _request(
        regulation_mappings=[
            RegulationMapping(
                law_name="AI 기본법",
                article_no="제27조",
                content="조항 본문",
                compliance="COMPLIANT",
                evidence="자율점검 기반 자동 매칭",
            )
        ]
    )

    result = compliance_service.generate_compliance_report(request)

    assert result.compliant_count == 1
    assert result.non_compliant_count == 0

    html = upload_sink[result.report_s3_key].decode("utf-8")
    assert "미준수로 판정된 조항 없음" in html


def test_omits_reference_section_without_audit_reference(patched_service):
    upload_sink, _ = patched_service

    result = compliance_service.generate_compliance_report(
        _request(audit_reference=None)
    )
    html = upload_sink[result.report_s3_key].decode("utf-8")

    assert "6. 참고 · 감사 결과 요약" not in html
    # 판정 본문은 그대로 있어야 한다
    assert "4. 조항별 준수 판정" in html
