"""리포트 Word 문서 생성에 공통으로 쓰는 도우미.

설명가능성(`report_docx.py`)·편향진단(`bias_report_docx.py`) 렌더러가 함께 쓴다.
용지·폰트·표·그림처럼 도메인과 무관한 것만 두고, 목차와 본문 구성은 각 렌더러가
담당한다.
"""

import base64
from io import BytesIO
from typing import Any, Iterable

from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.exceptions import (
    InvalidImageStreamError,
    UnexpectedEndOfFileError,
    UnrecognizedImageError,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT_NAME = "Malgun Gothic"


class DocxGenerationError(RuntimeError):
    """리포트의 Word 문서 생성에 실패한 경우."""


def format_number(
    value: float | int | None,
    digits: int = 4,
) -> str:
    if value is None:
        return "N/A"

    if isinstance(value, int):
        return str(value)

    return f"{value:.{digits}f}"


def format_list(values: list[str]) -> str:
    return ", ".join(values) if values else "없음"


def set_run_font(run, size: float | None = None, bold: bool | None = None):
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().rFonts.set(
        qn("w:eastAsia"),
        FONT_NAME,
    )

    if size is not None:
        run.font.size = Pt(size)

    if bold is not None:
        run.bold = bold


def configure_document(document: DocumentObject) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

    normal_style = document.styles["Normal"]
    normal_style.font.name = FONT_NAME
    normal_style._element.get_or_add_rPr().rFonts.set(
        qn("w:eastAsia"),
        FONT_NAME,
    )
    normal_style.font.size = Pt(9.5)

    for style_name, size in (
        ("Title", 22),
        ("Heading 1", 16),
        ("Heading 2", 12),
    ):
        style = document.styles[style_name]
        style.font.name = FONT_NAME
        style._element.get_or_add_rPr().rFonts.set(
            qn("w:eastAsia"),
            FONT_NAME,
        )
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(31, 41, 55)


def prevent_row_split(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    row_properties.append(cant_split)


def repeat_table_header(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    row_properties.append(table_header)


def set_cell_text(cell, value: Any, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(value))
    set_run_font(run, size=8.5, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(
    document: DocumentObject,
    headers: list[str],
    rows: Iterable[Iterable[Any]],
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True

    header_row = table.rows[0]
    repeat_table_header(header_row)
    prevent_row_split(header_row)

    for index, header in enumerate(headers):
        set_cell_text(header_row.cells[index], header, bold=True)

    for values in rows:
        row = table.add_row()
        prevent_row_split(row)

        for index, value in enumerate(values):
            set_cell_text(row.cells[index], value)

    document.add_paragraph()


def add_heading(
    document: DocumentObject,
    text: str,
    level: int,
) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(5)


def add_narrative(
    document: DocumentObject,
    text: str,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_after = Pt(6)

    run = paragraph.add_run(text)
    set_run_font(run, size=9.5)


def add_note(
    document: DocumentObject,
    text: str,
) -> None:
    """본문보다 작은 글씨의 보조 설명. HTML 리포트의 `.small` 문단에 대응한다."""

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_after = Pt(6)

    run = paragraph.add_run(text)
    set_run_font(run, size=8)


def add_bullets(
    document: DocumentObject,
    items: Iterable[str],
) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run(str(item))
        set_run_font(run, size=9.5)


def decode_figure(figure: dict[str, str]) -> BytesIO:
    data_uri = figure.get("data_uri", "")

    try:
        _, encoded = data_uri.split(",", maxsplit=1)
        return BytesIO(base64.b64decode(encoded, validate=True))
    except (ValueError, TypeError) as exception:
        raise DocxGenerationError(
            f"Word 리포트 그래프 데이터를 읽을 수 없습니다: "
            f"{figure.get('name', 'unknown')}"
        ) from exception


def add_figures(
    document: DocumentObject,
    figures: list[dict[str, str]],
    names: set[str] | None = None,
    section: str | None = None,
) -> None:
    """`names`(설명가능성) 또는 `section`(편향진단) 기준으로 그림을 삽입한다."""

    for figure in figures:
        if names is not None and figure.get("name") not in names:
            continue

        if section is not None and figure.get("section") != section:
            continue

        image_stream = decode_figure(figure)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_together = True

        run = paragraph.add_run()

        try:
            run.add_picture(image_stream, width=Cm(16.5))
        except (
            InvalidImageStreamError,
            UnexpectedEndOfFileError,
            UnrecognizedImageError,
            ValueError,
        ) as exception:
            raise DocxGenerationError(
                f"Word 리포트 그래프 형식이 올바르지 않습니다: "
                f"{figure.get('name', 'unknown')}"
            ) from exception

        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.keep_with_next = False

        caption_run = caption.add_run(figure.get("title", ""))
        set_run_font(caption_run, size=8)
        caption_run.italic = True
